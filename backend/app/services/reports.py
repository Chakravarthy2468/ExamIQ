import os
from datetime import datetime
from sqlalchemy.orm import Session
from app.db.models import Report, User, StudyPlan, TopicAnalytics, Course
from fpdf import FPDF
import openpyxl

REPORTS_DIR = "uploads/reports"
os.makedirs(REPORTS_DIR, exist_ok=True)

def generate_pdf_report(db: Session, user_id: int, course_id: int) -> Report:
    user = db.query(User).filter(User.id == user_id).first()
    course = db.query(Course).filter(Course.id == course_id).first()
    plan = db.query(StudyPlan).filter(StudyPlan.user_id == user_id, StudyPlan.course_id == course_id).order_by(StudyPlan.generated_at.desc()).first()
    
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=15)
    pdf.cell(200, 10, txt=f"ExamIQ Report - {course.name}", ln=True, align="C")
    
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt=f"Student: {user.full_name}", ln=True)
    pdf.cell(200, 10, txt=f"Date: {datetime.now().strftime('%Y-%m-%d')}", ln=True)
    
    readiness = plan.readiness_score if plan else 0.0
    pdf.cell(200, 10, txt=f"Overall Readiness Score: {readiness}%", ln=True)
    
    pdf.ln(10)
    pdf.set_font("Arial", 'I', size=10)
    pdf.multi_cell(0, 10, txt="Responsible AI Disclosure: AI answer evaluation is advisory and not equivalent to official faculty grading. AI-generated tutoring content and topic importance estimates may contain errors and do not guarantee future examination contents. Low-confidence evaluations may require human review.")
    
    filename = f"report_{user_id}_{course_id}_{int(datetime.now().timestamp())}.pdf"
    filepath = os.path.join(REPORTS_DIR, filename)
    pdf.output(filepath)
    
    report = Report(
        user_id=user_id,
        report_type="PDF",
        file_path=filepath,
        generated_at=datetime.now()
    )
    db.add(report)
    db.commit()
    db.refresh(report)
    return report

def generate_excel_report(db: Session, user_id: int, course_id: int) -> Report:
    user = db.query(User).filter(User.id == user_id).first()
    course = db.query(Course).filter(Course.id == course_id).first()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Analytics"
    
    ws.append(["ExamIQ Analytics Report"])
    ws.append(["Course", course.name])
    ws.append(["Student", user.full_name])
    ws.append([])
    ws.append(["Topic ID", "Frequency", "Total Marks", "Importance Score"])
    
    # Just grab some analytics
    analytics = db.query(TopicAnalytics).join(TopicAnalytics.topic).filter(TopicAnalytics.topic.has(unit=course.units)).all()
    for a in analytics:
        ws.append([a.topic_id, a.frequency, a.total_marks, a.importance_score])
        
    filename = f"report_{user_id}_{course_id}_{int(datetime.now().timestamp())}.xlsx"
    filepath = os.path.join(REPORTS_DIR, filename)
    wb.save(filepath)
    
    report = Report(
        user_id=user_id,
        report_type="EXCEL",
        file_path=filepath,
        generated_at=datetime.now()
    )
    db.add(report)
    db.commit()
    db.refresh(report)
    return report
